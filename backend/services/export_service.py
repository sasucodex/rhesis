import html
import os
import re
import tempfile
from typing import Optional
from pydantic import BaseModel
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from htmldocx import HtmlToDocx
from xhtml2pdf import pisa

class ExportRequest(BaseModel):
    text: str
    filename: str = "Trascrizione"
    date_str: Optional[str] = ""
    course_name: Optional[str] = ""
    professor_name: Optional[str] = ""

def sanitize_html_for_export(html_str: str) -> str:
    if not html_str:
        return ""
    text = re.sub(r"</?mark[^>]*>", "", html_str)
    text = re.sub(r"""<button[^>]*data-timestamp=["']([^"']+)["'][^>]*>.*?</button>""", r"[\1]", text)
    text = re.sub(r"<button[^>]*>(.*?)</button>", r"\1", text)
    text = re.sub(r"<h2>\s*</h2>", "", text)
    text = re.sub(r"<p>\s*</p>", "", text)
    return text

def generate_docx(req: ExportRequest) -> tuple[str, str]:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

        footer = section.footer
        footer_para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        footer_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        r1 = footer_para.add_run("Pagina ")
        r1.font.name = "Arial"
        r1.font.size = Pt(9)
        r1.font.color.rgb = RGBColor(113, 113, 122)

        fld_page = f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"><w:r><w:rPr><w:rFonts w:ascii="Arial"/><w:sz w:val="18"/><w:color w:val="71717A"/></w:rPr><w:t>1</w:t></w:r></w:fldSimple>'
        footer_para._p.append(parse_xml(fld_page))

        r2 = footer_para.add_run(" di ")
        r2.font.name = "Arial"
        r2.font.size = Pt(9)
        r2.font.color.rgb = RGBColor(113, 113, 122)

        fld_numpages = f'<w:fldSimple {nsdecls("w")} w:instr="NUMPAGES"><w:r><w:rPr><w:rFonts w:ascii="Arial"/><w:sz w:val="18"/><w:color w:val="71717A"/></w:rPr><w:t>1</w:t></w:r></w:fldSimple>'
        footer_para._p.append(parse_xml(fld_numpages))

    style_normal = doc.styles["Normal"]
    style_normal.font.name = "Arial"
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(24, 24, 27)

    title_text = req.filename.strip() if req.filename and req.filename.strip() else "Trascrizione"
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(title_text)
    title_run.font.name = "Arial"
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(9, 9, 11)
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(6)

    clean_date = req.date_str.strip() if req.date_str and req.date_str.strip() else ""
    clean_course = req.course_name.strip() if req.course_name and req.course_name.strip() else ""
    clean_prof = req.professor_name.strip() if req.professor_name and req.professor_name.strip() else ""

    meta_items = []
    if clean_date:
        meta_items.append(("Data", clean_date))
    if clean_course:
        meta_items.append(("Corso", clean_course))
    if clean_prof:
        meta_items.append(("Docente", clean_prof))

    for label, val in meta_items:
        mp = doc.add_paragraph()
        mp.paragraph_format.space_before = Pt(0)
        mp.paragraph_format.space_after = Pt(2)
        r_lbl = mp.add_run(f"{label.upper()}: ")
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(9)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = RGBColor(113, 113, 122)
        r_val = mp.add_run(val)
        r_val.font.name = "Arial"
        r_val.font.size = Pt(9.5)
        r_val.font.color.rgb = RGBColor(39, 39, 42)

    if meta_items:
        div_p = doc.add_paragraph()
        div_p.paragraph_format.space_before = Pt(4)
        div_p.paragraph_format.space_after = Pt(12)
    else:
        title_p.paragraph_format.space_after = Pt(14)

    clean_body = sanitize_html_for_export(req.text)
    parser = HtmlToDocx()
    parser.add_html_to_document(clean_body, doc)

    meta_labels = [m[0].upper() for m in meta_items]
    for p in doc.paragraphs:
        if p != title_p and not any(p.text.startswith(f"{lbl}:") for lbl in meta_labels):
            p.paragraph_format.line_spacing = 1.2
            if not p.text.startswith("##") and not p.style.name.startswith("Heading"):
                p.paragraph_format.space_after = Pt(6)

    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    temp_file.close()
    try:
        doc.save(temp_file.name)
    except Exception:
        if os.path.exists(temp_file.name):
            try:
                os.remove(temp_file.name)
            except OSError:
                pass
        raise
    download_name = f"{title_text}.docx"
    return temp_file.name, download_name

def generate_pdf(req: ExportRequest) -> tuple[str, str]:
    title_text = req.filename.strip() if req.filename and req.filename.strip() else "Trascrizione"
    clean_title = html.escape(title_text)
    clean_date = html.escape(req.date_str.strip()) if req.date_str and req.date_str.strip() else ""
    clean_course = html.escape(req.course_name.strip()) if req.course_name and req.course_name.strip() else ""
    clean_prof = html.escape(req.professor_name.strip()) if req.professor_name and req.professor_name.strip() else ""

    meta_rows = []
    if clean_date and clean_prof:
        meta_rows.append(f'<tr><td class="meta-val"><span class="meta-lbl">Data:</span> {clean_date}</td><td class="meta-val" style="text-align: right;"><span class="meta-lbl">Docente:</span> {clean_prof}</td></tr>')
    elif clean_date:
        meta_rows.append(f'<tr><td class="meta-val" colspan="2"><span class="meta-lbl">Data:</span> {clean_date}</td></tr>')
    elif clean_prof:
        meta_rows.append(f'<tr><td class="meta-val" colspan="2"><span class="meta-lbl">Docente:</span> {clean_prof}</td></tr>')

    if clean_course:
        meta_rows.append(f'<tr><td class="meta-val" colspan="2"><span class="meta-lbl">Corso:</span> {clean_course}</td></tr>')

    meta_html = ""
    if meta_rows:
        meta_html = f'<table class="meta-table">{"".join(meta_rows)}</table>'

    sanitized_body = sanitize_html_for_export(req.text)
    sanitized_body = re.sub(
        r'\[(\d{1,2}:\d{2}(?::\d{2})?)\]',
        r'<span style="color: #52525b; font-family: Courier, monospace; font-size: 9.5pt; font-weight: bold;">[\1]</span>',
        sanitized_body
    )

    full_html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
@page {{
    size: a4 portrait;
    margin-top: 20mm;
    margin-bottom: 22mm;
    margin-left: 20mm;
    margin-right: 20mm;
    @frame footer_frame {{
        -pdf-frame-content: footerContent;
        bottom: 8mm;
        margin-left: 20mm;
        margin-right: 20mm;
        height: 10mm;
    }}
}}
body {{
    font-family: Helvetica, Arial, sans-serif;
    font-size: 10.5pt;
    line-height: 1.55;
    color: #18181b;
}}
.header-box {{
    border-bottom: 1.5pt solid #18181b;
    padding-bottom: 10px;
    margin-bottom: 20px;
}}
.title {{
    font-size: 19pt;
    font-weight: bold;
    color: #09090b;
    margin-bottom: 6px;
}}
.meta-table {{
    width: 100%;
    margin-top: 6px;
    border-collapse: collapse;
}}
.meta-lbl {{
    font-size: 8.5pt;
    font-weight: bold;
    color: #71717a;
    text-transform: uppercase;
    letter-spacing: 0.5px;
}}
.meta-val {{
    font-size: 9.5pt;
    color: #27272a;
    padding-top: 2px;
    padding-bottom: 2px;
}}
.footer {{
    text-align: right;
    font-size: 8.5pt;
    color: #71717a;
    border-top: 0.5pt solid #e4e4e7;
    padding-top: 4px;
}}
h2 {{
    font-size: 13pt;
    font-weight: bold;
    color: #09090b;
    margin-top: 16px;
    margin-bottom: 6px;
    border-bottom: 0.5pt solid #e4e4e7;
    padding-bottom: 3px;
}}
p {{
    margin-bottom: 10px;
    text-align: justify;
}}
</style>
</head>
<body>
<div id="footerContent" class="footer">
    Pagina <pdf:pagenumber> di <pdf:pagecount>
</div>

<div class="header-box">
    <div class="title">{clean_title}</div>
    {meta_html}
</div>

{sanitized_body}
</body>
</html>"""

    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
    temp_file.close()
    try:
        with open(temp_file.name, "wb") as f:
            res = pisa.CreatePDF(full_html, dest=f)
            if res.err:
                raise RuntimeError("Errore durante la generazione del file PDF")
    except Exception:
        if os.path.exists(temp_file.name):
            try:
                os.remove(temp_file.name)
            except OSError:
                pass
        raise

    download_name = f"{title_text}.pdf"
    return temp_file.name, download_name
