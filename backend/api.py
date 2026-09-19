from fastapi import FastAPI, UploadFile, File, HTTPException, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel
from typing import Optional
import os
import re
import time
import html
from google import genai
from google.genai import types, errors
import httpx
import json
import tempfile
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import shutil
import markdown
import uuid
from htmldocx import HtmlToDocx
from xhtml2pdf import pisa
import mimetypes

try:
    from backend.database import get_db_connection, init_db
except ImportError:
    from database import get_db_connection, init_db

app = FastAPI(title="Rhesis Transcription Server")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)
CONFIG_DIR = os.path.expanduser("~/.config/rhesis")
CONFIG_FILE = os.path.join(CONFIG_DIR, "config.json")
UPLOAD_DIR = os.path.join(CONFIG_DIR, "uploads")
os.makedirs(CONFIG_DIR, exist_ok=True)
os.makedirs(UPLOAD_DIR, exist_ok=True)

def clean_transcription_markdown(raw_text: str) -> str:
    if not raw_text:
        return ""
    text = raw_text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r'(?m)^[ \t]*##[ \t]*\n+(?:[ \t]*\n)*[ \t]*(\[\d{1,2}:\d{2}(?::\d{2})?\]\s*[^\n]+)', r'## \1', text)
    text = re.sub(r'(?m)^[ \t]*##[ \t]*\n+(?:[ \t]*\n)*[ \t]*([A-Za-zÀ-ÿ0-9][^\n]+)', r'## \1', text)
    text = re.sub(r'(?m)^[ \t]*(\[\d{1,2}:\d{2}(?::\d{2})?\])\s*##\s*', r'## \1 ', text)
    text = re.sub(r'(?m)^[ \t]*(\[\d{1,2}:\d{2}(?::\d{2})?\])\s*\n+(?:[ \t]*\n)*##\s*', r'## \1 ', text)
    text = re.sub(r'(?m)^[ \t]*##\s*$', '', text)

    def clean_heading_and_following(match):
        ts = match.group(1)
        title = match.group(2).strip()
        sep = match.group(3)
        following = match.group(4)
        if re.match(r'\[\d{1,2}:\d{2}(?::\d{2})?\]', following.strip()):
            return f"## {title}{sep}{following}"
        return f"## {title}{sep}{ts} {following.lstrip()}"

    text = re.sub(
        r'(?m)^[ \t]*##\s*(\[\d{1,2}:\d{2}(?::\d{2})?\])\s*([^\n]+)(\n+(?:[ \t]*\n)*)([^\n]+)',
        clean_heading_and_following,
        text
    )
    text = re.sub(r'(?m)^[ \t]*##\s*\[\d{1,2}:\d{2}(?::\d{2})?\]\s*([^\n]+)', r'## \1', text)
    return text

def normalize_transcript_html(html: str) -> str:
    if not html:
        return ""
    html = re.sub(r'<p>\s*(\[\d{1,2}:\d{2}(?::\d{2})?\])\s*##\s*(.*?)</p>', r'<h2>\2</h2>', html)
    html = re.sub(r'<p>\s*##\s*(\[\d{1,2}:\d{2}(?::\d{2})?\]\s*.*?)</p>', r'<h2>\1</h2>', html)
    html = re.sub(r'<p>\s*##\s*(.*?)</p>', r'<h2>\1</h2>', html)
    html = re.sub(r'<h2>\s*</h2>\s*<p>(\s*\[\d{1,2}:\d{2}(?::\d{2})?\].*?)</p>', r'<h2>\1</h2>', html)
    html = re.sub(r'<h2>\s*</h2>\s*', '', html)

    def clean_h2_and_p(match):
        h2_full = match.group(1)
        h2_inner = match.group(2)
        sep = match.group(3)
        p_open = match.group(4)
        p_inner = match.group(5)

        ts_match = re.search(r'\[\d{1,2}:\d{2}(?::\d{2})?\]', h2_inner)
        if not ts_match:
            return match.group(0)

        ts = ts_match.group(0)
        clean_h2_inner = re.sub(r'\[\d{1,2}:\d{2}(?::\d{2})?\]\s*', '', h2_inner).strip()
        clean_h2 = f"<h2>{clean_h2_inner}</h2>"

        if re.match(r'^\s*\[\d{1,2}:\d{2}(?::\d{2})?\]', p_inner):
            return f"{clean_h2}{sep}{p_open}{p_inner}"
        else:
            return f"{clean_h2}{sep}{p_open}{ts} {p_inner.lstrip()}"

    html = re.sub(
        r'(<h2>(.*?)</h2>)(\s*)(<p[^>]*>)(.*?)(?=</p>|\Z)',
        clean_h2_and_p,
        html,
        flags=re.DOTALL
    )

    def strip_ts_from_h2(match):
        inner = match.group(1)
        clean_inner = re.sub(r'\[\d{1,2}:\d{2}(?::\d{2})?\]\s*', '', inner).strip()
        return f"<h2>{clean_inner}</h2>"

    html = re.sub(r'<h2>(.*?)</h2>', strip_ts_from_h2, html, flags=re.DOTALL)
    return html

init_db()
def get_saved_api_key():
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, 'r') as f:
                config = json.load(f)
                return config.get("api_key")
        except:
            pass
    return None
def save_api_key(key: str):
    os.makedirs(CONFIG_DIR, exist_ok=True)
    with open(CONFIG_FILE, 'w') as f:
        json.dump({"api_key": key}, f)
class SetupRequest(BaseModel):
    api_key: str
class CourseCreateRequest(BaseModel):
    name: str
    color: Optional[str] = "#2563eb"
    professor_name: Optional[str] = None
class CourseUpdateRequest(BaseModel):
    name: Optional[str] = None
    color: Optional[str] = None
    professor_name: Optional[str] = None
class TranscriptionResponse(BaseModel):
    id: int
    transcript: str
    status: str
    audio_preserved: Optional[bool] = False
class ExportRequest(BaseModel):
    text: str
    filename: str = "Trascrizione"
    date_str: Optional[str] = ""
    course_name: Optional[str] = ""
    professor_name: Optional[str] = ""
class UpdateTranscriptRequest(BaseModel):
    transcript: Optional[str] = None
    filename: Optional[str] = None
    course_id: Optional[int] = None
    course_name: Optional[str] = None
    professor_name: Optional[str] = None
def check_api_key_status(key: str, retry_on_unreachable: bool = True) -> tuple[Optional[bool], str]:
    if not key or not isinstance(key, str) or len(key.strip()) < 10:
        return False, "invalid"

    attempts = 2 if retry_on_unreachable else 1
    cleaned_key = key.strip()

    for attempt in range(attempts):
        try:
            client = genai.Client(api_key=cleaned_key)
            client.models.get(model="gemini-3-flash-preview")
            return True, "valid"
        except errors.ClientError as e:
            if getattr(e, "code", None) in (400, 401, 403):
                return False, "invalid"
            return None, "unreachable"
        except errors.ServerError:
            if attempt < attempts - 1:
                time.sleep(1)
                continue
            return None, "unreachable"
        except (httpx.RequestError, httpx.TimeoutException, OSError):
            if attempt < attempts - 1:
                time.sleep(1)
                continue
            return None, "unreachable"
        except Exception:
            return None, "unreachable"

    return None, "unreachable"

def verify_api_key(key: str) -> bool:
    is_valid, _ = check_api_key_status(key)
    return is_valid is True

@app.get("/")
def read_root():
    return {"message": "Rhesis Server is running! 🚀"}

@app.get("/status")
def check_status():
    api_key = get_saved_api_key()
    is_configured = bool(api_key and api_key.strip())
    if not is_configured:
        return {
            "server_running": True,
            "api_key_configured": False,
            "api_key_valid": False,
            "api_key_status": "unconfigured"
        }

    is_valid, status = check_api_key_status(api_key, retry_on_unreachable=True)
    return {
        "server_running": True, 
        "api_key_configured": True,
        "api_key_valid": is_valid,
        "api_key_status": status
    }

@app.post("/setup")
def setup_api_key(req: SetupRequest):
    key = req.api_key.strip() if req.api_key else ""
    if not key:
        raise HTTPException(status_code=400, detail="Il codice di accesso non può essere vuoto")

    is_valid, status = check_api_key_status(key, retry_on_unreachable=False)
    if status == "invalid":
        raise HTTPException(status_code=400, detail="Il codice inserito non sembra corretto o è incompleto. Assicurati di averlo copiato per intero e riprova.")
    if status == "unreachable":
        raise HTTPException(status_code=400, detail="Impossibile verificare il codice con Google: connessione internet assente o irraggiungibile.")

    save_api_key(key)
    return {"message": "Configurazione salvata con successo"}

@app.delete("/setup")
def delete_api_key():
    if os.path.exists(CONFIG_FILE):
        os.remove(CONFIG_FILE)
    return {"message": "Codice di accesso rimosso con successo"}
import threading

TASKS = {}
tasks_lock = threading.Lock()

def update_task(task_id: str, status: str, progress: int, message: str, result=None, error=None):
    with tasks_lock:
        if task_id not in TASKS:
            TASKS[task_id] = {
                "task_id": task_id,
                "created_at": time.time(),
            }
        TASKS[task_id].update({
            "status": status,
            "progress": progress,
            "message": message,
            "result": result,
            "error": error,
            "updated_at": time.time(),
        })

def cleanup_old_tasks():
    with tasks_lock:
        now = time.time()
        expired = [tid for tid, t in TASKS.items() if now - t.get("created_at", now) > 7200]
        for tid in expired:
            del TASKS[tid]

@app.get("/courses")
def get_courses():
    with get_db_connection() as conn:
        c = conn.cursor()
        c.execute("""
            SELECT c.id, c.name, c.color, c.professor_name, c.created_at,
                   COUNT(t.id) as transcriptions_count
            FROM courses c
            LEFT JOIN transcriptions t ON t.course_id = c.id AND t.status = 'success'
            GROUP BY c.id
            ORDER BY c.name COLLATE NOCASE ASC
        """)
        rows = c.fetchall()
        return [dict(r) for r in rows]

@app.post("/courses")
def create_course(req: CourseCreateRequest):
    name = req.name.strip() if req.name else ""
    if not name:
        raise HTTPException(status_code=400, detail="Il nome del corso è obbligatorio.")
    color = req.color.strip() if req.color else "#2563eb"
    prof = req.professor_name.strip() if req.professor_name else None
    with get_db_connection() as conn:
        c = conn.cursor()
        c.execute("SELECT id FROM courses WHERE LOWER(name) = LOWER(?)", (name,))
        if c.fetchone():
            raise HTTPException(status_code=400, detail="Un corso con questo nome esiste già.")
        c.execute(
            "INSERT INTO courses (name, color, professor_name) VALUES (?, ?, ?)",
            (name, color, prof)
        )
        course_id = c.lastrowid
        conn.commit()
        return {
            "id": course_id,
            "name": name,
            "color": color,
            "professor_name": prof,
            "transcriptions_count": 0
        }

@app.put("/courses/{course_id}")
def update_course(course_id: int, req: CourseUpdateRequest):
    with get_db_connection() as conn:
        c = conn.cursor()
        c.execute("SELECT id, name, color, professor_name FROM courses WHERE id = ?", (course_id,))
        course = c.fetchone()
        if not course:
            raise HTTPException(status_code=404, detail="Corso non trovato.")

        new_name = req.name.strip() if req.name is not None else course["name"]
        if not new_name:
            raise HTTPException(status_code=400, detail="Il nome del corso non può essere vuoto.")

        if new_name.lower() != course["name"].lower():
            c.execute("SELECT id FROM courses WHERE LOWER(name) = LOWER(?) AND id != ?", (new_name, course_id))
            if c.fetchone():
                raise HTTPException(status_code=400, detail="Un altro corso con questo nome esiste già.")

        new_color = req.color.strip() if req.color is not None else course["color"]
        new_prof = req.professor_name.strip() if req.professor_name is not None else course["professor_name"]
        if new_prof == "":
            new_prof = None

        c.execute(
            "UPDATE courses SET name = ?, color = ?, professor_name = ? WHERE id = ?",
            (new_name, new_color, new_prof, course_id)
        )
        c.execute("UPDATE transcriptions SET course_name = ? WHERE course_id = ?", (new_name, course_id))
        conn.commit()

        c.execute("""
            SELECT c.id, c.name, c.color, c.professor_name, c.created_at,
                   COUNT(t.id) as transcriptions_count
            FROM courses c
            LEFT JOIN transcriptions t ON t.course_id = c.id AND t.status = 'success'
            WHERE c.id = ?
            GROUP BY c.id
        """, (course_id,))
        updated = c.fetchone()
        return dict(updated)

@app.delete("/courses/{course_id}")
def delete_course(course_id: int):
    with get_db_connection() as conn:
        c = conn.cursor()
        c.execute("UPDATE transcriptions SET course_id = NULL WHERE course_id = ?", (course_id,))
        c.execute("DELETE FROM courses WHERE id = ?", (course_id,))
        conn.commit()
        return {"message": "Corso eliminato con successo"}

@app.get("/history")
def get_history(course_id: Optional[int] = None):
    with get_db_connection() as conn:
        c = conn.cursor()
        query = """
            SELECT t.id, t.filename, t.file_path, t.status, t.created_at, t.transcript, t.audio_preserved,
                   t.course_id,
                   COALESCE(c.name, t.course_name) as course_name,
                   COALESCE(c.professor_name, t.professor_name) as professor_name,
                   c.color as course_color
            FROM transcriptions t
            LEFT JOIN courses c ON t.course_id = c.id
            WHERE t.status = 'success'
        """
        params = []
        if course_id is not None and course_id > 0:
            query += " AND t.course_id = ?"
            params.append(course_id)
        query += " ORDER BY t.id DESC"
        c.execute(query, tuple(params))
        rows = c.fetchall()
    result = []
    for row in rows:
        item = dict(row)
        file_path = item.pop("file_path", None)
        item["audio_preserved"] = bool(item.get("audio_preserved", 0)) and bool(file_path and os.path.exists(file_path))
        result.append(item)
    return result

def build_fts_query(user_query: str) -> str:
    cleaned = user_query.strip()
    if not cleaned:
        return ""
    tokens = re.findall(r'[^\W_]+', cleaned, re.UNICODE)
    if not tokens:
        return ""
    joined = " ".join(tokens)
    return f'"{joined}"*'

def focus_snippet_around_mark(snippet: str, max_chars_before: int = 25) -> str:
    if not snippet or "<mark>" not in snippet:
        return snippet.strip() if snippet else ""
    first_mark_idx = snippet.find("<mark>")
    if first_mark_idx <= max_chars_before:
        return snippet.strip()
    pre_text = snippet[:first_mark_idx]
    cut_pos = first_mark_idx - max_chars_before
    space_pos = pre_text.find(" ", cut_pos)
    if space_pos != -1 and space_pos < first_mark_idx:
        trimmed_pre = "..." + pre_text[space_pos:].lstrip()
    else:
        trimmed_pre = "..." + pre_text[cut_pos:]
    return (trimmed_pre + snippet[first_mark_idx:]).strip()

@app.get("/search")
def search_transcriptions(q: str = ""):
    fts_q = build_fts_query(q)
    if not fts_q:
        return []
    with get_db_connection() as conn:
        c = conn.cursor()
        try:
            c.execute("""
                SELECT t.id, t.filename, t.file_path, t.transcript, t.status, t.audio_preserved,
                       t.course_id,
                       COALESCE(c.name, t.course_name) as course_name,
                       COALESCE(c.professor_name, t.professor_name) as professor_name,
                       c.color as course_color,
                       t.created_at,
                       snippet(transcriptions_fts, 1, '<mark>', '</mark>', '...', 12) as content_snippet,
                       snippet(transcriptions_fts, 0, '<mark>', '</mark>', '...', 12) as title_snippet
                FROM transcriptions_fts f
                JOIN transcriptions t ON t.id = f.rowid
                LEFT JOIN courses c ON t.course_id = c.id
                WHERE transcriptions_fts MATCH ? AND t.status = 'success'
                ORDER BY rank
            """, (fts_q,))
            rows = c.fetchall()
        except Exception:
            pattern = f"%{q.strip()}%"
            c.execute("""
                SELECT t.id, t.filename, t.file_path, t.transcript, t.status, t.audio_preserved,
                       t.course_id,
                       COALESCE(c.name, t.course_name) as course_name,
                       COALESCE(c.professor_name, t.professor_name) as professor_name,
                       c.color as course_color,
                       t.created_at,
                       '' as content_snippet, '' as title_snippet
                FROM transcriptions t
                LEFT JOIN courses c ON t.course_id = c.id
                WHERE t.status = 'success' AND (t.filename LIKE ? OR t.transcript LIKE ?)
                ORDER BY t.id DESC
            """, (pattern, pattern))
            rows = c.fetchall()

    result = []
    for row in rows:
        item = dict(row)
        file_path = item.pop("file_path", None)
        item["audio_preserved"] = bool(item.get("audio_preserved", 0)) and bool(file_path and os.path.exists(file_path))
        c_snip = item.pop("content_snippet", "") or ""
        t_snip = item.pop("title_snippet", "") or ""
        if "<mark>" in c_snip:
            item["snippet"] = focus_snippet_around_mark(c_snip)
        elif "<mark>" in t_snip:
            item["snippet"] = focus_snippet_around_mark(t_snip)
        else:
            item["snippet"] = (c_snip.strip() or t_snip.strip())[:80]
        result.append(item)
    return result


def process_transcription_core(
    api_key: str,
    file_path: str,
    record_id: int,
    enable_chapters: bool = False,
    enable_timestamps: bool = False,
    preserve_audio: bool = False,
    task_id: Optional[str] = None
):
    uploaded_file = None
    client = None
    try:
        client = genai.Client(api_key=api_key)
        if task_id:
            update_task(task_id, "upload_google", 35, "Caricamento audio su Google AI Studio in corso...")
        uploaded_file = client.files.upload(file=file_path)

        if task_id:
            update_task(task_id, "processing_google", 55, "I server Google stanno elaborando il file audio...")
        file_info = client.files.get(name=uploaded_file.name)
        poll_count = 0
        while file_info.state.name == 'PROCESSING':
            time.sleep(2)
            poll_count += 1
            if task_id:
                prog = min(55 + poll_count * 2, 70)
                update_task(task_id, "processing_google", prog, "I server Google stanno elaborando il file audio...")
            file_info = client.files.get(name=uploaded_file.name)

        if file_info.state.name == 'FAILED':
            raise Exception("Google backend failed processing audio.")

        if task_id:
            update_task(task_id, "gemini_generating", 75, "Trascrizione in corso con Gemini...")

        system_instruction = (
            "Sei un trascrittore professionale e accademico. Il tuo compito è produrre una trascrizione testuale PAROLA PER PAROLA (verbatim) dell'audio fornito. "
            "REGOLA FONDAMENTALE 1: NON riassumere MAI, NON saltare argomenti e NON tagliare il testo. Devi trascrivere ogni singola frase detta, parola per parola. "
            "REGOLA FONDAMENTALE 2: Non inserire mai frasi di cortesia o saluti (es. 'Ecco la trascrizione'). Inizia direttamente con il contenuto. "
            "Fai massima attenzione all'esattezza dei termini tecnici, scientifici e alla corretta capitalizzazione."
        )
        prompt = "Esegui la trascrizione parola per parola dell'intero audio, senza tralasciare nulla."
        if enable_chapters:
            prompt += " Durante la trascrizione fedele, dividi il discorso inserendo dei titoli di capitolo (usa il formato Markdown ## per i titoli) quando il professore cambia argomento. IMPORTANTE: i capitoli servono solo a separare visivamente il testo, NON devi assolutamente riassumere il contenuto sotto di essi. Continua a trascrivere parola per parola."
        else:
            prompt += " Non inserire titoli di capitolo, trascrivi tutto in un flusso continuo diviso semplicemente in paragrafi."
        if enable_timestamps:
            prompt += " Inserisci i timestamp nel formato [MM:SS] all'inizio di ogni cambio logico di argomento o blocco di discorso."
        else:
            prompt += " ASSOLUTAMENTE NON INSERIRE timestamp o minutaggi nel testo."
        if enable_chapters and enable_timestamps:
            prompt += (
                " NOTA BENE PER I TITOLI DI CAPITOLO E I TIMESTAMP: quando introduci un capitolo, scrivi tassativamente '## Titolo del capitolo' sulla stessa riga (es. '## Introduzione e riepilogo'). "
                "Non lasciare mai i cancelletti '##' vuoti o isolati sulla riga. "
                "REGOLA TASSATIVA: nei titoli di capitolo '##' NON inserire MAI timestamp o minutaggi [MM:SS]. "
                "Il timestamp [MM:SS] deve comparire ESCLUSIVAMENTE all'inizio del paragrafo di testo sottostante, mai dentro o prima del titolo del capitolo (Esempio corretto:\n"
                "## Titolo del capitolo\n\n[00:00] Testo della lezione...)."
            )
        elif enable_chapters:
            prompt += " NOTA BENE PER I TITOLI DI CAPITOLO: quando introduci un capitolo, scrivi tassativamente '## Titolo del capitolo' sulla stessa riga. Non lasciare mai cancelletti '##' vuoti o isolati."

        max_retries = 3
        response = None
        for attempt in range(max_retries):
            reconnect_timer = None
            try:
                if attempt > 0 and task_id:
                    update_task(
                        task_id,
                        "gemini_generating",
                        75,
                        f"Nuovo tentativo ({attempt + 1}/{max_retries}) in corso con Gemini..."
                    )
                    reconnect_timer = threading.Timer(
                        2.0,
                        update_task,
                        args=(task_id, "gemini_generating", 75, "Trascrizione in corso con Gemini...")
                    )
                    reconnect_timer.daemon = True
                    reconnect_timer.start()

                response = client.models.generate_content(
                    model="gemini-3-flash-preview",
                    contents=[uploaded_file, prompt],
                    config=types.GenerateContentConfig(
                        system_instruction=system_instruction,
                        temperature=0.2
                    )
                )
                break
            except Exception as gen_err:
                err_str = str(gen_err)
                if ('503' in err_str or 'high demand' in err_str.lower() or 'unavailable' in err_str.lower()) and attempt < max_retries - 1:
                    backoff = (2 ** attempt) * 2
                    for rem in range(backoff, 0, -1):
                        if task_id:
                            update_task(
                                task_id,
                                "gemini_generating",
                                75,
                                f"Server Google occupati, nuovo tentativo tra {rem}s..."
                            )
                        time.sleep(1)
                else:
                    raise gen_err
            finally:
                if reconnect_timer:
                    reconnect_timer.cancel()

        raw_text = response.text
        cleaned_text = clean_transcription_markdown(raw_text)
        transcript_html = normalize_transcript_html(markdown.markdown(cleaned_text))
        if uploaded_file:
            try:
                client.files.delete(name=uploaded_file.name)
            except Exception:
                pass

        with get_db_connection() as conn:
            c = conn.cursor()
            c.execute("UPDATE transcriptions SET transcript = ?, status = 'success', created_at = CURRENT_TIMESTAMP WHERE id = ?", (transcript_html, record_id))
            conn.commit()

        if not preserve_audio and os.path.exists(file_path):
            try:
                os.remove(file_path)
            except Exception:
                pass

        if task_id:
            update_task(
                task_id,
                "completed",
                100,
                "Trascrizione completata con successo!",
                result={
                    "id": record_id,
                    "transcript": transcript_html,
                    "status": "success",
                    "audio_preserved": preserve_audio
                }
            )
        return transcript_html
    except Exception as e:
        error_msg = str(e)
        is_network_err = any(kw in error_msg.lower() for kw in [
            'temporary failure in name resolution',
            'name or service not known',
            'nodename nor servname provided',
            'getaddrinfo failed',
            'failed to establish a new connection',
            'network is unreachable',
            'network unreachable',
            'connection refused',
            'errno -3',
            'errno -2',
            'errno 101',
            'connecterror',
            'connecttimeout',
            'networkerror',
            'socket.gaierror',
            'connection error',
            'connection reset',
            'max retries exceeded with url',
            'failed to resolve',
            'no address associated with hostname'
        ])
        if is_network_err:
            error_msg = "Connessione internet assente o non raggiungibile. Verifica la tua connessione Wi-Fi o di rete e riprova."
        elif any(kw in error_msg.lower() for kw in ['401', 'unauthenticated', 'api key not valid', 'api_key_invalid', 'permission_denied', 'consumer_invalid']):
            error_msg = "Codice di accesso non valido o scaduto. Per avviare la trascrizione è necessario configurare un codice personale funzionante nelle Impostazioni."
        elif '503' in error_msg or 'high demand' in error_msg.lower():
            error_msg = "I server di Google sono momentaneamente saturi (High Demand). Riprova tra qualche minuto."

        if uploaded_file and client:
            try:
                client.files.delete(name=uploaded_file.name)
            except Exception:
                pass

        with get_db_connection() as conn:
            c = conn.cursor()
            c.execute("DELETE FROM transcriptions WHERE id = ?", (record_id,))
            conn.commit()

        if os.path.exists(file_path):
            try:
                os.remove(file_path)
            except Exception:
                pass

        if task_id:
            update_task(task_id, "error", 0, error_msg, error=error_msg)

        raise Exception(error_msg)

def transcription_worker(task_id: str, api_key: str, file_path: str, record_id: int, enable_chapters: bool, enable_timestamps: bool, preserve_audio: bool):
    try:
        process_transcription_core(
            api_key=api_key,
            file_path=file_path,
            record_id=record_id,
            enable_chapters=enable_chapters,
            enable_timestamps=enable_timestamps,
            preserve_audio=preserve_audio,
            task_id=task_id
        )
    except Exception:
        pass

@app.post("/transcribe/")
def transcribe(
    file: UploadFile = File(...),
    enable_chapters: bool = Form(False),
    enable_timestamps: bool = Form(False),
    preserve_audio: bool = Form(False),
    course_id: Optional[int] = Form(None),
    course_name: Optional[str] = Form(None),
    professor_name: Optional[str] = Form(None)
):
    cleanup_old_tasks()
    api_key = get_saved_api_key()
    if not api_key or len(api_key.strip()) < 10:
        raise HTTPException(status_code=401, detail="Codice di accesso non configurato o scaduto. Configuralo nelle Impostazioni.")
    ext = os.path.splitext(file.filename)[1].lower() if file.filename else ""
    allowed_exts = {'.mp3', '.m4a', '.wav', '.ogg', '.flac', '.aac', '.mp4', '.webm', '.mpeg', '.mpga', '.amr'}
    if ext not in allowed_exts:
        raise HTTPException(status_code=400, detail=f"Formato file non supportato ({ext}). Usa un formato audio valido (es. MP3, M4A, WAV).")
    safe_filename = f"{int(time.time())}_{uuid.uuid4().hex[:8]}{ext}"
    file_path = os.path.join(UPLOAD_DIR, safe_filename)
    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
    with get_db_connection() as conn:
        c = conn.cursor()

        course_id_val = None
        if course_id is not None and course_id > 0:
            course_id_val = course_id
            c.execute("SELECT name, professor_name FROM courses WHERE id = ?", (course_id_val,))
            c_row = c.fetchone()
            if c_row:
                if not course_name:
                    course_name = c_row[0]
                if not professor_name and c_row[1]:
                    professor_name = c_row[1]

        c.execute(
            "INSERT INTO transcriptions (filename, file_path, status, audio_preserved, course_name, professor_name, course_id) VALUES (?, ?, 'processing', ?, ?, ?, ?)",
            (file.filename, file_path, 1 if preserve_audio else 0, course_name, professor_name, course_id_val)
        )
        record_id = c.lastrowid
        conn.commit()

    task_id = uuid.uuid4().hex
    with tasks_lock:
        TASKS[task_id] = {
            "task_id": task_id,
            "record_id": record_id,
            "status": "upload_local",
            "progress": 15,
            "message": "File audio salvato in locale. Inizializzazione pipeline...",
            "result": None,
            "error": None,
            "created_at": time.time(),
            "updated_at": time.time()
        }

    worker = threading.Thread(
        target=transcription_worker,
        args=(task_id, api_key, file_path, record_id, enable_chapters, enable_timestamps, preserve_audio),
        daemon=True
    )
    worker.start()

    return {
        "task_id": task_id,
        "id": record_id,
        "record_id": record_id,
        "status": "upload_local"
    }

@app.get("/task/{task_id}")
def get_task_status(task_id: str):
    with tasks_lock:
        task = TASKS.get(task_id)
        if not task:
            raise HTTPException(status_code=404, detail="Task non trovato")
        return dict(task)

@app.api_route("/audio/{record_id}", methods=["GET", "HEAD"])
def stream_audio(record_id: int):
    with get_db_connection() as conn:
        c = conn.cursor()
        c.execute("SELECT file_path, audio_preserved, filename FROM transcriptions WHERE id = ?", (record_id,))
        row = c.fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="Trascrizione non trovata")
    file_path, audio_preserved, orig_filename = row[0], bool(row[1]), row[2]
    if not audio_preserved or not file_path or not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail="File audio non disponibile o non conservato per questa trascrizione")
    
    media_type, _ = mimetypes.guess_type(file_path)
    if not media_type or not media_type.startswith("audio/"):
        ext = os.path.splitext(file_path)[1].lower()
        audio_mimes = {
            ".mp3": "audio/mpeg",
            ".m4a": "audio/mp4",
            ".wav": "audio/wav",
            ".ogg": "audio/ogg",
            ".flac": "audio/flac",
            ".aac": "audio/aac",
            ".webm": "audio/webm",
            ".mp4": "audio/mp4",
            ".mpeg": "audio/mpeg",
            ".mpga": "audio/mpeg",
            ".amr": "audio/amr",
        }
        media_type = audio_mimes.get(ext, "audio/mpeg")
    
    return FileResponse(
        file_path,
        media_type=media_type.lower() if media_type else "audio/mpeg",
        filename=orig_filename,
        content_disposition_type="inline",
        headers={"Accept-Ranges": "bytes"}
    )

@app.put("/transcript/{record_id}")
def update_transcript(record_id: int, req: UpdateTranscriptRequest):
    with get_db_connection() as conn:
        c = conn.cursor()
        updates = []
        params = []
        if req.transcript is not None:
            updates.append("transcript = ?")
            params.append(normalize_transcript_html(req.transcript))
        if req.filename is not None:
            updates.append("filename = ?")
            params.append(req.filename)
        if req.course_id is not None:
            if req.course_id > 0:
                updates.append("course_id = ?")
                params.append(req.course_id)
                c.execute("SELECT name, professor_name FROM courses WHERE id = ?", (req.course_id,))
                c_row = c.fetchone()
                if c_row:
                    updates.append("course_name = ?")
                    params.append(c_row[0])
                    if c_row[1] and req.professor_name is None:
                        updates.append("professor_name = ?")
                        params.append(c_row[1])
            else:
                updates.append("course_id = NULL")
                updates.append("course_name = NULL")
                updates.append("professor_name = NULL")
        if req.course_name is not None and (req.course_id is None or req.course_id > 0):
            updates.append("course_name = ?")
            params.append(req.course_name)
        if req.professor_name is not None and (req.course_id is None or req.course_id > 0):
            updates.append("professor_name = ?")
            params.append(req.professor_name)
        if updates:
            params.append(record_id)
            c.execute(f"UPDATE transcriptions SET {', '.join(updates)} WHERE id = ?", tuple(params))
            conn.commit()
    return {"message": "Success"}

@app.delete("/transcript/{record_id}")
def delete_transcript(record_id: int):
    with get_db_connection() as conn:
        c = conn.cursor()
        c.execute("SELECT file_path FROM transcriptions WHERE id = ?", (record_id,))
        row = c.fetchone()
        if row and row[0] and os.path.exists(row[0]):
            try:
                os.remove(row[0])
            except:
                pass
        c.execute("DELETE FROM transcriptions WHERE id = ?", (record_id,))
        conn.commit()
    return {"message": "Success"}

def sanitize_html_for_export(html_str: str) -> str:
    if not html_str:
        return ""
    text = re.sub(r"</?mark[^>]*>", "", html_str)
    text = re.sub(r"""<button[^>]*data-timestamp=["']([^"']+)["'][^>]*>.*?</button>""", r"[\1]", text)
    text = re.sub(r"<button[^>]*>(.*?)</button>", r"\1", text)
    text = re.sub(r"<h2>\s*</h2>", "", text)
    text = re.sub(r"<p>\s*</p>", "", text)
    return text

@app.post("/export/word")
def export_word(req: ExportRequest):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

        footer = section.footer
        footer_para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        footer_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        r1 = footer_para.add_run("Pagina ")
        r1.font.name = "Arial"
        r1.font.size = Pt(9)
        r1.font.color.rgb = RGBColor(113, 113, 122)

        fld_page = f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"><w:r><w:rPr><w:rFonts w:ascii="Arial"/><w:sz w:val="18"/><w:color w:val="71717A"/></w:rPr><w:t>1</w:t></w:r></w:fldSimple>'
        footer_para._p.append(parse_xml(fld_page))

        r2 = footer_para.add_run(" di ")
        r2.font.name = "Arial"
        r2.font.size = Pt(9)
        r2.font.color.rgb = RGBColor(113, 113, 122)

        fld_numpages = f'<w:fldSimple {nsdecls("w")} w:instr="NUMPAGES"><w:r><w:rPr><w:rFonts w:ascii="Arial"/><w:sz w:val="18"/><w:color w:val="71717A"/></w:rPr><w:t>1</w:t></w:r></w:fldSimple>'
        footer_para._p.append(parse_xml(fld_numpages))

    style_normal = doc.styles["Normal"]
    style_normal.font.name = "Arial"
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(24, 24, 27)

    title_text = (req.filename or "Trascrizione").strip()
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(title_text)
    title_run.font.name = "Arial"
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(9, 9, 11)
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(6)

    clean_date = req.date_str.strip() if req.date_str and req.date_str.strip() else ""
    clean_course = req.course_name.strip() if req.course_name and req.course_name.strip() else ""
    clean_prof = req.professor_name.strip() if req.professor_name and req.professor_name.strip() else ""

    meta_items = []
    if clean_date:
        meta_items.append(("Data", clean_date))
    if clean_course:
        meta_items.append(("Corso", clean_course))
    if clean_prof:
        meta_items.append(("Docente", clean_prof))

    for label, val in meta_items:
        mp = doc.add_paragraph()
        mp.paragraph_format.space_before = Pt(0)
        mp.paragraph_format.space_after = Pt(2)
        r_lbl = mp.add_run(f"{label.upper()}: ")
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(9)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = RGBColor(113, 113, 122)
        r_val = mp.add_run(val)
        r_val.font.name = "Arial"
        r_val.font.size = Pt(9.5)
        r_val.font.color.rgb = RGBColor(39, 39, 42)

    if meta_items:
        div_p = doc.add_paragraph()
        div_p.paragraph_format.space_before = Pt(4)
        div_p.paragraph_format.space_after = Pt(12)
    else:
        title_p.paragraph_format.space_after = Pt(14)

    clean_body = sanitize_html_for_export(req.text)
    parser = HtmlToDocx()
    parser.add_html_to_document(clean_body, doc)

    meta_labels = [m[0].upper() for m in meta_items]
    for p in doc.paragraphs:
        if p != title_p and not any(p.text.startswith(f"{lbl}:") for lbl in meta_labels):
            p.paragraph_format.line_spacing = 1.2
            if not p.text.startswith("##") and not p.style.name.startswith("Heading"):
                p.paragraph_format.space_after = Pt(6)

    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    doc.save(temp_file.name)
    download_name = f"{title_text}.docx"
    return FileResponse(
        temp_file.name,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=download_name
    )

@app.post("/export/pdf")
def export_pdf(req: ExportRequest):
    title_text = (req.filename or "Trascrizione").strip()
    clean_title = html.escape(title_text)
    clean_date = html.escape(req.date_str.strip()) if req.date_str and req.date_str.strip() else ""
    clean_course = html.escape(req.course_name.strip()) if req.course_name and req.course_name.strip() else ""
    clean_prof = html.escape(req.professor_name.strip()) if req.professor_name and req.professor_name.strip() else ""

    meta_rows = []
    if clean_date and clean_prof:
        meta_rows.append(f'<tr><td class="meta-val"><span class="meta-lbl">Data:</span> {clean_date}</td><td class="meta-val" style="text-align: right;"><span class="meta-lbl">Docente:</span> {clean_prof}</td></tr>')
    elif clean_date:
        meta_rows.append(f'<tr><td class="meta-val" colspan="2"><span class="meta-lbl">Data:</span> {clean_date}</td></tr>')
    elif clean_prof:
        meta_rows.append(f'<tr><td class="meta-val" colspan="2"><span class="meta-lbl">Docente:</span> {clean_prof}</td></tr>')

    if clean_course:
        meta_rows.append(f'<tr><td class="meta-val" colspan="2"><span class="meta-lbl">Corso:</span> {clean_course}</td></tr>')

    meta_html = ""
    if meta_rows:
        meta_html = f'<table class="meta-table">{"".join(meta_rows)}</table>'

    sanitized_body = sanitize_html_for_export(req.text)
    sanitized_body = re.sub(
        r'\[(\d{1,2}:\d{2}(?::\d{2})?)\]',
        r'<span style="color: #52525b; font-family: Courier, monospace; font-size: 9.5pt; font-weight: bold;">[\1]</span>',
        sanitized_body
    )

    full_html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
@page {{
    size: a4 portrait;
    margin-top: 20mm;
    margin-bottom: 22mm;
    margin-left: 20mm;
    margin-right: 20mm;
    @frame footer_frame {{
        -pdf-frame-content: footerContent;
        bottom: 8mm;
        margin-left: 20mm;
        margin-right: 20mm;
        height: 10mm;
    }}
}}
body {{
    font-family: Helvetica, Arial, sans-serif;
    font-size: 10.5pt;
    line-height: 1.55;
    color: #18181b;
}}
.header-box {{
    border-bottom: 1.5pt solid #18181b;
    padding-bottom: 10px;
    margin-bottom: 20px;
}}
.title {{
    font-size: 19pt;
    font-weight: bold;
    color: #09090b;
    margin-bottom: 6px;
}}
.meta-table {{
    width: 100%;
    margin-top: 6px;
    border-collapse: collapse;
}}
.meta-lbl {{
    font-size: 8.5pt;
    font-weight: bold;
    color: #71717a;
    text-transform: uppercase;
    letter-spacing: 0.5px;
}}
.meta-val {{
    font-size: 9.5pt;
    color: #27272a;
    padding-top: 2px;
    padding-bottom: 2px;
}}
.footer {{
    text-align: right;
    font-size: 8.5pt;
    color: #71717a;
    border-top: 0.5pt solid #e4e4e7;
    padding-top: 4px;
}}
h2 {{
    font-size: 13pt;
    font-weight: bold;
    color: #09090b;
    margin-top: 16px;
    margin-bottom: 6px;
    border-bottom: 0.5pt solid #e4e4e7;
    padding-bottom: 3px;
}}
p {{
    margin-bottom: 10px;
    text-align: justify;
}}
</style>
</head>
<body>
<div id="footerContent" class="footer">
    Pagina <pdf:pagenumber> di <pdf:pagecount>
</div>

<div class="header-box">
    <div class="title">{clean_title}</div>
    {meta_html}
</div>

{sanitized_body}
</body>
</html>"""

    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
    with open(temp_file.name, "wb") as f:
        res = pisa.CreatePDF(full_html, dest=f)
        if res.err:
            raise HTTPException(status_code=500, detail="Errore durante la generazione del file PDF")

    download_name = f"{title_text}.pdf"
    return FileResponse(
        temp_file.name,
        media_type="application/pdf",
        filename=download_name
    )
if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=8000)
