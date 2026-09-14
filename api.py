from fastapi import FastAPI, UploadFile, File, HTTPException, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel
from typing import Optional
import os
import re
import time
from google import genai
from google.genai import types
import json
import tempfile
import sqlite3
from datetime import datetime
from docx import Document
from fpdf import FPDF
import shutil
import markdown
import uuid
from htmldocx import HtmlToDocx
import mimetypes

app = FastAPI(title="Rhesis Transcription Server")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)
CONFIG_DIR = os.path.expanduser("~/.config/rhesis")
CONFIG_FILE = os.path.join(CONFIG_DIR, "config.json")
DB_FILE = os.path.join(CONFIG_DIR, "database.db")
UPLOAD_DIR = os.path.join(CONFIG_DIR, "uploads")
os.makedirs(CONFIG_DIR, exist_ok=True)
os.makedirs(UPLOAD_DIR, exist_ok=True)

def init_db():
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    c.execute('''
        CREATE TABLE IF NOT EXISTS transcriptions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            filename TEXT NOT NULL,
            file_path TEXT NOT NULL,
            transcript TEXT,
            status TEXT NOT NULL,
            audio_preserved BOOLEAN DEFAULT 0,
            created_at DATETIME DEFAULT CURRENT_TIMESTAMP
        )
    ''')
    c.execute("PRAGMA table_info(transcriptions)")
    columns = [row[1] for row in c.fetchall()]
    if "audio_preserved" not in columns:
        c.execute("ALTER TABLE transcriptions ADD COLUMN audio_preserved BOOLEAN DEFAULT 0")
    conn.commit()
    conn.close()

init_db()
def get_saved_api_key():
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, 'r') as f:
                config = json.load(f)
                return config.get("api_key")
        except:
            pass
    return None
def save_api_key(key: str):
    os.makedirs(CONFIG_DIR, exist_ok=True)
    with open(CONFIG_FILE, 'w') as f:
        json.dump({"api_key": key}, f)
class SetupRequest(BaseModel):
    api_key: str
class TranscriptionResponse(BaseModel):
    id: int
    transcript: str
    status: str
    audio_preserved: Optional[bool] = False
from docx.shared import Pt, RGBColor
class ExportRequest(BaseModel):
    text: str
    filename: str = "Trascrizione"
    date_str: str = ""
class UpdateTranscriptRequest(BaseModel):
    transcript: Optional[str] = None
    filename: Optional[str] = None
def verify_api_key(key: str) -> bool:
    if not key or not isinstance(key, str) or len(key.strip()) < 10:
        return False
    try:
        client = genai.Client(api_key=key.strip())
        client.models.get(model="gemini-3-flash-preview")
        return True
    except Exception:
        return False

@app.get("/")
def read_root():
    return {"message": "Rhesis Server is running! 🚀"}

@app.get("/status")
def check_status():
    api_key = get_saved_api_key()
    is_configured = bool(api_key and api_key.strip())
    is_valid = verify_api_key(api_key) if is_configured else False
    return {
        "server_running": True, 
        "api_key_configured": is_configured,
        "api_key_valid": is_valid
    }

@app.post("/setup")
def setup_api_key(req: SetupRequest):
    key = req.api_key.strip() if req.api_key else ""
    if not verify_api_key(key):
        raise HTTPException(status_code=400, detail="Chiave API Google non valida o revocata")
    save_api_key(key)
    return {"message": "Configurazione salvata con successo"}

@app.delete("/setup")
def delete_api_key():
    if os.path.exists(CONFIG_FILE):
        os.remove(CONFIG_FILE)
    return {"message": "API Key rimossa con successo"}
import threading

TASKS = {}
tasks_lock = threading.Lock()

def update_task(task_id: str, status: str, progress: int, message: str, result=None, error=None):
    with tasks_lock:
        if task_id not in TASKS:
            TASKS[task_id] = {
                "task_id": task_id,
                "created_at": time.time(),
            }
        TASKS[task_id].update({
            "status": status,
            "progress": progress,
            "message": message,
            "result": result,
            "error": error,
            "updated_at": time.time(),
        })

def cleanup_old_tasks():
    with tasks_lock:
        now = time.time()
        expired = [tid for tid, t in TASKS.items() if now - t.get("created_at", now) > 7200]
        for tid in expired:
            del TASKS[tid]

@app.get("/history")
def get_history():
    conn = sqlite3.connect(DB_FILE)
    conn.row_factory = sqlite3.Row
    c = conn.cursor()
    c.execute("SELECT id, filename, file_path, status, created_at, transcript, audio_preserved FROM transcriptions WHERE status = 'success' ORDER BY id DESC")
    rows = c.fetchall()
    conn.close()
    result = []
    for row in rows:
        item = dict(row)
        file_path = item.pop("file_path", None)
        item["audio_preserved"] = bool(item.get("audio_preserved", 0)) and bool(file_path and os.path.exists(file_path))
        result.append(item)
    return result

def process_transcription_core(
    api_key: str,
    file_path: str,
    record_id: int,
    enable_chapters: bool = False,
    enable_timestamps: bool = False,
    preserve_audio: bool = False,
    task_id: Optional[str] = None
):
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    uploaded_file = None
    client = None
    try:
        client = genai.Client(api_key=api_key)
        if task_id:
            update_task(task_id, "upload_google", 35, "Caricamento audio su Google AI Studio in corso...")
        uploaded_file = client.files.upload(file=file_path)

        if task_id:
            update_task(task_id, "processing_google", 55, "I server Google stanno elaborando il file audio...")
        file_info = client.files.get(name=uploaded_file.name)
        poll_count = 0
        while file_info.state.name == 'PROCESSING':
            time.sleep(2)
            poll_count += 1
            if task_id:
                prog = min(55 + poll_count * 2, 70)
                update_task(task_id, "processing_google", prog, "I server Google stanno elaborando il file audio...")
            file_info = client.files.get(name=uploaded_file.name)

        if file_info.state.name == 'FAILED':
            raise Exception("Google backend failed processing audio.")

        if task_id:
            update_task(task_id, "gemini_generating", 75, "Trascrizione in corso con Gemini...")

        system_instruction = (
            "Sei un trascrittore professionale e accademico. Il tuo compito è produrre una trascrizione testuale PAROLA PER PAROLA (verbatim) dell'audio fornito. "
            "REGOLA FONDAMENTALE 1: NON riassumere MAI, NON saltare argomenti e NON tagliare il testo. Devi trascrivere ogni singola frase detta, parola per parola. "
            "REGOLA FONDAMENTALE 2: Non inserire mai frasi di cortesia o saluti (es. 'Ecco la trascrizione'). Inizia direttamente con il contenuto. "
            "Fai massima attenzione all'esattezza dei termini tecnici, scientifici e alla corretta capitalizzazione."
        )
        prompt = "Esegui la trascrizione parola per parola dell'intero audio, senza tralasciare nulla."
        if enable_chapters:
            prompt += " Durante la trascrizione fedele, dividi il discorso inserendo dei titoli di capitolo (usa il formato Markdown ## per i titoli) quando il professore cambia argomento. IMPORTANTE: i capitoli servono solo a separare visivamente il testo, NON devi assolutamente riassumere il contenuto sotto di essi. Continua a trascrivere parola per parola."
        else:
            prompt += " Non inserire titoli di capitolo, trascrivi tutto in un flusso continuo diviso semplicemente in paragrafi."
        if enable_timestamps:
            prompt += " Inserisci i timestamp nel formato [MM:SS] all'inizio di ogni cambio logico di argomento o blocco di discorso."
        else:
            prompt += " ASSOLUTAMENTE NON INSERIRE timestamp o minutaggi nel testo."

        max_retries = 3
        response = None
        for attempt in range(max_retries):
            reconnect_timer = None
            try:
                if attempt > 0 and task_id:
                    update_task(
                        task_id,
                        "gemini_generating",
                        75,
                        f"Nuovo tentativo ({attempt + 1}/{max_retries}) in corso con Gemini..."
                    )
                    reconnect_timer = threading.Timer(
                        2.0,
                        update_task,
                        args=(task_id, "gemini_generating", 75, "Trascrizione in corso con Gemini...")
                    )
                    reconnect_timer.daemon = True
                    reconnect_timer.start()

                response = client.models.generate_content(
                    model="gemini-3-flash-preview",
                    contents=[uploaded_file, prompt],
                    config=types.GenerateContentConfig(
                        system_instruction=system_instruction,
                        temperature=0.2
                    )
                )
                break
            except Exception as gen_err:
                err_str = str(gen_err)
                if ('503' in err_str or 'high demand' in err_str.lower() or 'unavailable' in err_str.lower()) and attempt < max_retries - 1:
                    backoff = (2 ** attempt) * 2
                    for rem in range(backoff, 0, -1):
                        if task_id:
                            update_task(
                                task_id,
                                "gemini_generating",
                                75,
                                f"Server Google occupati, nuovo tentativo tra {rem}s..."
                            )
                        time.sleep(1)
                else:
                    raise gen_err
            finally:
                if reconnect_timer:
                    reconnect_timer.cancel()

        raw_text = response.text
        cleaned_text = re.sub(r'(?m)^(\s*\[\d{1,2}:\d{2}(?::\d{2})?\])\s*##\s*', r'## \1 ', raw_text)
        cleaned_text = re.sub(r'(\[\d{1,2}:\d{2}(?::\d{2})?\])\s*##\s*', r'\1 ', cleaned_text)
        transcript_html = markdown.markdown(cleaned_text)
        if uploaded_file:
            try:
                client.files.delete(name=uploaded_file.name)
            except Exception:
                pass

        c.execute("UPDATE transcriptions SET transcript = ?, status = 'success', created_at = CURRENT_TIMESTAMP WHERE id = ?", (transcript_html, record_id))
        conn.commit()

        if not preserve_audio and os.path.exists(file_path):
            try:
                os.remove(file_path)
            except Exception:
                pass

        if task_id:
            update_task(
                task_id,
                "completed",
                100,
                "Trascrizione completata con successo!",
                result={
                    "id": record_id,
                    "transcript": transcript_html,
                    "status": "success",
                    "audio_preserved": preserve_audio
                }
            )
        return transcript_html
    except Exception as e:
        error_msg = str(e)
        if '503' in error_msg or 'high demand' in error_msg.lower():
            error_msg = "I server di Google sono momentaneamente saturi (High Demand). Riprova tra qualche minuto."

        if uploaded_file and client:
            try:
                client.files.delete(name=uploaded_file.name)
            except Exception:
                pass

        c.execute("DELETE FROM transcriptions WHERE id = ?", (record_id,))
        conn.commit()

        if os.path.exists(file_path):
            try:
                os.remove(file_path)
            except Exception:
                pass

        if task_id:
            update_task(task_id, "error", 0, error_msg, error=error_msg)

        raise Exception(error_msg)
    finally:
        conn.close()

def transcription_worker(task_id: str, api_key: str, file_path: str, record_id: int, enable_chapters: bool, enable_timestamps: bool, preserve_audio: bool):
    try:
        process_transcription_core(
            api_key=api_key,
            file_path=file_path,
            record_id=record_id,
            enable_chapters=enable_chapters,
            enable_timestamps=enable_timestamps,
            preserve_audio=preserve_audio,
            task_id=task_id
        )
    except Exception:
        pass

@app.post("/transcribe/")
def transcribe(
    file: UploadFile = File(...),
    enable_chapters: bool = Form(False),
    enable_timestamps: bool = Form(False),
    preserve_audio: bool = Form(False)
):
    cleanup_old_tasks()
    api_key = get_saved_api_key()
    if not api_key:
        raise HTTPException(status_code=401, detail="API Key not configured.")
    ext = os.path.splitext(file.filename)[1].lower() if file.filename else ""
    allowed_exts = {'.mp3', '.m4a', '.wav', '.ogg', '.flac', '.aac', '.mp4', '.webm', '.mpeg', '.mpga', '.amr'}
    if ext not in allowed_exts:
        raise HTTPException(status_code=400, detail=f"Formato file non supportato ({ext}). Usa un formato audio valido (es. MP3, M4A, WAV).")
    safe_filename = f"{int(time.time())}_{uuid.uuid4().hex[:8]}{ext}"
    file_path = os.path.join(UPLOAD_DIR, safe_filename)
    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    c.execute("INSERT INTO transcriptions (filename, file_path, status, audio_preserved) VALUES (?, ?, 'processing', ?)", (file.filename, file_path, 1 if preserve_audio else 0))
    record_id = c.lastrowid
    conn.commit()
    conn.close()

    task_id = uuid.uuid4().hex
    with tasks_lock:
        TASKS[task_id] = {
            "task_id": task_id,
            "record_id": record_id,
            "status": "upload_local",
            "progress": 15,
            "message": "File audio salvato in locale. Inizializzazione pipeline...",
            "result": None,
            "error": None,
            "created_at": time.time(),
            "updated_at": time.time()
        }

    worker = threading.Thread(
        target=transcription_worker,
        args=(task_id, api_key, file_path, record_id, enable_chapters, enable_timestamps, preserve_audio),
        daemon=True
    )
    worker.start()

    return {
        "task_id": task_id,
        "id": record_id,
        "record_id": record_id,
        "status": "upload_local"
    }

@app.get("/task/{task_id}")
def get_task_status(task_id: str):
    with tasks_lock:
        task = TASKS.get(task_id)
        if not task:
            raise HTTPException(status_code=404, detail="Task non trovato")
        return dict(task)

@app.api_route("/audio/{record_id}", methods=["GET", "HEAD"])
def stream_audio(record_id: int):
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    c.execute("SELECT file_path, audio_preserved, filename FROM transcriptions WHERE id = ?", (record_id,))
    row = c.fetchone()
    conn.close()
    if not row:
        raise HTTPException(status_code=404, detail="Trascrizione non trovata")
    file_path, audio_preserved, orig_filename = row[0], bool(row[1]), row[2]
    if not audio_preserved or not file_path or not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail="File audio non disponibile o non conservato per questa trascrizione")
    
    media_type, _ = mimetypes.guess_type(file_path)
    if not media_type or not media_type.startswith("audio/"):
        ext = os.path.splitext(file_path)[1].lower()
        audio_mimes = {
            ".mp3": "audio/mpeg",
            ".m4a": "audio/mp4",
            ".wav": "audio/wav",
            ".ogg": "audio/ogg",
            ".flac": "audio/flac",
            ".aac": "audio/aac",
            ".webm": "audio/webm",
            ".mp4": "audio/mp4",
            ".mpeg": "audio/mpeg",
            ".mpga": "audio/mpeg",
            ".amr": "audio/amr",
        }
        media_type = audio_mimes.get(ext, "audio/mpeg")
    
    return FileResponse(
        file_path,
        media_type=media_type.lower() if media_type else "audio/mpeg",
        filename=orig_filename,
        content_disposition_type="inline",
        headers={"Accept-Ranges": "bytes"}
    )
@app.put("/transcript/{record_id}")
def update_transcript(record_id: int, req: UpdateTranscriptRequest):
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    updates = []
    params = []
    if req.transcript is not None:
        updates.append("transcript = ?")
        params.append(req.transcript)
    if req.filename is not None:
        updates.append("filename = ?")
        params.append(req.filename)
    if updates:
        params.append(record_id)
        c.execute(f"UPDATE transcriptions SET {', '.join(updates)} WHERE id = ?", tuple(params))
        conn.commit()
    conn.close()
    return {"message": "Success"}
@app.delete("/transcript/{record_id}")
def delete_transcript(record_id: int):
    conn = sqlite3.connect(DB_FILE)
    c = conn.cursor()
    c.execute("SELECT file_path FROM transcriptions WHERE id = ?", (record_id,))
    row = c.fetchone()
    if row and row[0] and os.path.exists(row[0]):
        try:
            os.remove(row[0])
        except:
            pass
    c.execute("DELETE FROM transcriptions WHERE id = ?", (record_id,))
    conn.commit()
    conn.close()
    return {"message": "Success"}
@app.post("/export/word")
def export_word(req: ExportRequest):
    doc = Document()
    heading = doc.add_heading(req.filename, 0)
    if req.date_str:
        run = heading.add_run(f"  {req.date_str}")
        run.font.color.rgb = RGBColor(128, 128, 128)
        run.font.size = Pt(14)
    new_parser = HtmlToDocx()
    new_parser.add_html_to_document(req.text, doc)
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    doc.save(temp_file.name)
    return FileResponse(temp_file.name, media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document', filename="Trascrizione.docx")
from xhtml2pdf import pisa
@app.post("/export/pdf")
def export_pdf(req: ExportRequest):
    date_html = f' <span style="color: gray; font-size: 14pt;">{req.date_str}</span>' if req.date_str else ""
    html_content = f"""
    <html>
    <head>
    <meta charset="utf-8">
    <style>
    body {  font-family: Helvetica, sans-serif; font-size: 12pt; color: #333; } 
    h1 {  text-align: left; font-size: 18pt; margin-bottom: 20px; border-bottom: 1px solid #eee; padding-bottom: 10px; } 
    </style>
    </head>
    <body>
    <h1>{req.filename}{date_html}</h1>
    {req.text}
    </body>
    </html>
    """
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
    with open(temp_file.name, "wb") as f:
        pisa.CreatePDF(html_content, dest=f)
    return FileResponse(temp_file.name, media_type='application/pdf', filename="Trascrizione.pdf")
if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=8000)
